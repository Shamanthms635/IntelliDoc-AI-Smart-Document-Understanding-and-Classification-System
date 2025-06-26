import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import os
from pathlib import Path
from datetime import datetime
import pickle
from docx import Document
from transformers import pipeline
import pdfplumber

# Load model and vectorizer
with open("document_model.pkl", "rb") as f:
    model = pickle.load(f)
with open("vectorizer.pkl", "rb") as f:
    vectorizer = pickle.load(f)

# Load summarizer
summarizer = pipeline("summarization", model="sshleifer/distilbart-cnn-12-6")

# Extract text from file
def extract_text(file_path):
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    elif suffix == ".pdf":
        text = ""
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text
    elif suffix == ".docx":
        doc = Document(str(file_path))
        return "\n".join([p.text for p in doc.paragraphs])
    return None

# Process file and display results
def process_file(file_path):
    
    result_text = ("")
    file_path = Path(file_path)
    if not file_path.exists():
        messagebox.showerror("Error", "File not found!")
        return

    text = extract_text(file_path)
    if not text or not text.strip():
        messagebox.showwarning("Warning", "No text extracted from file.")
        return

    summary = summarizer(text[:1024], max_length=80, min_length=20, do_sample=False)[0]['summary_text']
    vec = vectorizer.transform([text])
    prediction = model.predict(vec)[0]

    meta = f"\nFile: {file_path.name}\nSize: {round(file_path.stat().st_size / 1024, 2)} KB\nModified: {datetime.fromtimestamp(file_path.stat().st_mtime)}"
    meta += f"\n\n📝 Summary:\n{summary}"
    meta += f"\n\n🧠 Predicted Document Type: {prediction.capitalize()}"

    text_box.delete("1.0", tk.END)
    text_box.insert(tk.END, meta)

# File/folder dialog handlers
def browse_file():
    filepath = filedialog.askopenfilename(filetypes=[("Documents", "*.pdf *.docx *.txt")])
    if filepath:
        process_file(filepath)

def browse_folder():
    folderpath = filedialog.askdirectory()
    if folderpath:
        supported_ext = [".pdf", ".docx", ".txt"]
        files = [f for f in Path(folderpath).iterdir() if f.suffix in supported_ext and f.is_file()]
        if not files:
            messagebox.showinfo("Info", "No supported documents found in the folder.")
            return
        all_results = ""
        for f in files:
            text = extract_text(f)
            if not text:
                continue
            summary = summarizer(text[:1024], max_length=80, min_length=20, do_sample=False)[0]['summary_text']
            vec = vectorizer.transform([text])
            prediction = model.predict(vec)[0]
            all_results += f"\n{'='*60}\n{f.name}\nSize: {round(f.stat().st_size/1024, 2)} KB\nModified: {datetime.fromtimestamp(f.stat().st_mtime)}"
            all_results += f"\nSummary: {summary}"
            all_results += f"\nPrediction: {prediction.capitalize()}\n"

        text_box.delete("1.0", tk.END)
        text_box.insert(tk.END, all_results)

# GUI Setup
root = tk.Tk()
root.title("📄 Smart Document Classifier")
root.geometry("800x600")
root.configure(bg="white")

# Welcome
title = tk.Label(root, text="Smart Document Classifier", font=("cambria", 20, "bold"), bg="white", fg="blue")
title.pack(pady=10)

# Buttons
btn_frame = tk.Frame(root, bg="white")
btn_frame.pack(pady=10)

file_btn = ttk.Button(btn_frame, text="Browse File", command=browse_file)
file_btn.grid(row=0, column=0, padx=10)

folder_btn = ttk.Button(btn_frame, text="Browse Folder", command=browse_folder)
folder_btn.grid(row=0, column=1, padx=10)

# Scrollable Text Box
text_frame = tk.Frame(root, bg="white")
text_frame.pack(fill="both", expand=True, padx=20, pady=10)

scrollbar = tk.Scrollbar(text_frame)
scrollbar.pack(side="right", fill="y")

text_box = tk.Text(text_frame, wrap="word", yscrollcommand=scrollbar.set, font=("Helvetica", 11))
text_box.pack(fill="both", expand=True)

scrollbar.config(command=text_box.yview)

root.mainloop()
